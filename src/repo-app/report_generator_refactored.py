"""
Excel report generation service (refactored for CLEAN code principles).

Author: Automated Software Engineering Team
Date: February 2026
"""

from pathlib import Path
from typing import List

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

from .models import TestMetrics


# Constants for Excel formatting
HEADER_COLOR = "366092"
HEADER_FONT_COLOR = "FFFFFF"
SUCCESS_COLOR = "C6EFCE"
FAILED_COLOR = "FFC7CE"
MAX_COLUMN_WIDTH = 50


class ExcelReportGenerator:
    """
    Generate Excel reports from test metrics.
    
    Follows Single Responsibility Principle - only handles Excel generation.
    """
    
    def generate_report(self, metrics_list: List[TestMetrics], output_file: Path):
        """
        Generate formatted Excel report.
        
        Args:
            metrics_list: List of TestMetrics objects
            output_file: Path to output Excel file
        """
        print(f"\n📊 Generating Excel report: {output_file}")
        
        # Convert to DataFrame
        df = self._prepare_dataframe(metrics_list)
        
        # Create Excel workbook with formatting
        wb = self._create_formatted_workbook(df)
        
        # Save workbook
        wb.save(output_file)
        print(f"✅ Report saved: {output_file}")
    
    def _prepare_dataframe(self, metrics_list: List[TestMetrics]) -> pd.DataFrame:
        """Convert metrics to pandas DataFrame with proper formatting"""
        data = [m.to_dict() for m in metrics_list]
        df = pd.DataFrame(data)
        
        # Convert list fields to strings for Excel
        df['test_frameworks'] = df['test_frameworks'].apply(self._join_list)
        df['languages'] = df['languages'].apply(self._join_list)
        df['cicd_workflows'] = df['cicd_workflows'].apply(self._join_list)
        
        # Reorder columns
        df = df[self._get_column_order()]
        
        return df
    
    def _join_list(self, value) -> str:
        """Join list values for Excel display"""
        return ', '.join(value) if value else ''
    
    def _get_column_order(self) -> List[str]:
        """Define column order for report"""
        return [
            'repo_name', 'analysis_status', 'last_analyzed',
            'unit_test_count', 'unit_test_coverage_pct', 'unit_test_lines_covered', 'unit_test_lines_total',
            'feature_test_scenarios', 'feature_test_files',
            'performance_test_count', 'performance_test_files',
            'e2e_test_count', 'e2e_test_files',
            'smoke_test_count', 'smoke_test_files',
            'total_test_files', 'test_frameworks', 'languages',
            'total_commits', 'commits_last_30_days', 'commits_last_90_days', 'commits_last_year',
            'avg_commits_per_week', 'active_contributors',
            'has_cicd_pipeline', 'cicd_platform', 'has_automated_testing',
            'has_security_scanning', 'has_deployment_automation', 'cicd_workflows',
            'repo_url', 'error_message'
        ]
    
    def _create_formatted_workbook(self, df: pd.DataFrame) -> Workbook:
        """Create Excel workbook with formatting"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Repository Test Analysis"
        
        # Write header
        self._write_header(ws)
        
        # Write data
        self._write_data(ws, df)
        
        # Format worksheet
        self._format_worksheet(ws)
        
        return wb
    
    def _write_header(self, ws):
        """Write formatted header row"""
        headers = self._get_header_labels()
        header_style = self._get_header_style()
        
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_style["fill"]
            cell.font = header_style["font"]
            cell.alignment = header_style["alignment"]
    
    def _write_data(self, ws, df: pd.DataFrame):
        """Write data rows with conditional formatting"""
        for row_idx, row_data in enumerate(dataframe_to_rows(df, index=False, header=False), start=2):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                
                # Apply conditional formatting
                self._apply_cell_formatting(cell, col_idx, value)
    
    def _apply_cell_formatting(self, cell, col_idx: int, value):
        """Apply conditional formatting to cells"""
        # Coverage percentage format
        if col_idx == 5:  # Coverage %
            cell.number_format = '0.00'
        
        # Status color coding
        if col_idx == 2:  # Status column
            if value == "success":
                cell.fill = PatternFill(
                    start_color=SUCCESS_COLOR,
                    end_color=SUCCESS_COLOR,
                    fill_type="solid"
                )
            elif value == "failed":
                cell.fill = PatternFill(
                    start_color=FAILED_COLOR,
                    end_color=FAILED_COLOR,
                    fill_type="solid"
                )
    
    def _format_worksheet(self, ws):
        """Apply worksheet-level formatting"""
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = self._calculate_column_width(column)
            column_letter = column[0].column_letter
            ws.column_dimensions[column_letter].width = max_length
        
        # Freeze header row
        ws.freeze_panes = "A2"
    
    def _calculate_column_width(self, column) -> int:
        """Calculate optimal column width"""
        max_length = 0
        
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        
        return min(max_length + 2, MAX_COLUMN_WIDTH)
    
    def _get_header_labels(self) -> List[str]:
        """Get human-readable header labels"""
        return [
            'Repository', 'Status', 'Last Analyzed',
            'Unit Tests', 'Coverage %', 'Lines Covered', 'Total Lines',
            'Feature Scenarios', 'Feature Files',
            'Perf Tests', 'Perf Files',
            'E2E Tests', 'E2E Files',
            'Smoke Tests', 'Smoke Files',
            'Total Test Files', 'Frameworks', 'Languages',
            'Total Commits', 'Commits (30d)', 'Commits (90d)', 'Commits (1y)',
            'Commits/Week', 'Contributors',
            'Has CI/CD', 'CI/CD Platform', 'Auto Testing',
            'Security Scan', 'Auto Deploy', 'Workflows',
            'Repository URL', 'Error Message'
        ]
    
    def _get_header_style(self) -> dict:
        """Get header row style"""
        return {
            "fill": PatternFill(
                start_color=HEADER_COLOR,
                end_color=HEADER_COLOR,
                fill_type="solid"
            ),
            "font": Font(
                bold=True,
                color=HEADER_FONT_COLOR,
                size=11
            ),
            "alignment": Alignment(
                horizontal='center',
                vertical='center'
            )
        }
