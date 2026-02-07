"""
Excel report generator for repository test analysis.

Author: Automated Software Engineering Team
Date: February 2026
"""

from pathlib import Path
from typing import List

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

from models import TestMetrics


class ExcelReportGenerator:
    """Generate Excel report from test metrics"""
    
    def generate_report(self, metrics_list: List[TestMetrics], output_file: Path):
        """Generate Excel report with formatting"""
        print(f"\n📊 Generating Excel report: {output_file}")
        
        # Convert to DataFrame
        data = [m.to_dict() for m in metrics_list]
        df = pd.DataFrame(data)
        
        # Convert lists to strings for Excel
        df['test_frameworks'] = df['test_frameworks'].apply(lambda x: ', '.join(x) if x else '')
        df['languages'] = df['languages'].apply(lambda x: ', '.join(x) if x else '')
        df['cicd_workflows'] = df['cicd_workflows'].apply(lambda x: ', '.join(x) if x else '')
        
        # Reorder columns
        column_order = [
            'repo_name', 'analysis_status', 'last_analyzed',
            'unit_test_count', 'unit_test_coverage_pct', 'unit_test_lines_covered', 'unit_test_lines_total',
            'feature_test_scenarios', 'feature_test_files',
            'performance_test_count', 'performance_test_files',
            'e2e_test_count', 'e2e_test_files',
            'smoke_test_count', 'smoke_test_files',
            'total_test_files', 'test_frameworks', 'languages',
            'total_commits', 'commits_last_30_days', 'commits_last_90_days', 'commits_last_year',
            'avg_commits_per_week', 'active_contributors',
            'has_cicd_pipeline', 'cicd_platform', 'has_automated_testing',
            'has_security_scanning', 'has_deployment_automation', 'cicd_workflows',
            'repo_url', 'error_message'
        ]
        df = df[column_order]
        
        # Create Excel workbook with formatting
        wb = Workbook()
        ws = wb.active
        ws.title = "Repository Test Analysis"
        
        # Write header with formatting
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        
        headers = [
            'Repository', 'Status', 'Last Analyzed',
            'Unit Tests', 'Coverage %', 'Lines Covered', 'Total Lines',
            'Feature Scenarios', 'Feature Files',
            'Perf Tests', 'Perf Files',
            'E2E Tests', 'E2E Files',
            'Smoke Tests', 'Smoke Files',
            'Total Test Files', 'Frameworks', 'Languages',
            'Total Commits', 'Commits (30d)', 'Commits (90d)', 'Commits (1y)',
            'Commits/Week', 'Contributors',
            'Has CI/CD', 'CI/CD Platform', 'Auto Testing',
            'Security Scan', 'Auto Deploy', 'Workflows',
            'Repository URL', 'Error Message'
        ]
        
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Write data
        for row_idx, row_data in enumerate(dataframe_to_rows(df, index=False, header=False), start=2):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                
                # Format percentage
                if col_idx == 5:  # Coverage %
                    cell.number_format = '0.00'
                
                # Status color coding
                if col_idx == 2:  # Status
                    if value == "success":
                        cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                    elif value == "failed":
                        cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except (TypeError, AttributeError):
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Freeze header row
        ws.freeze_panes = "A2"
        
        # Save workbook
        wb.save(output_file)
        print(f"✅ Report saved: {output_file}")
